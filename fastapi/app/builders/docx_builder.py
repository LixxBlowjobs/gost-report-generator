from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from io import BytesIO
from typing import Dict, Any, List, Optional
import os, re

TEMPLATE_PATH = "/app/data/template.docx"
BOOKMARK_ID = 0


def _next_bm_id():
    global BOOKMARK_ID
    BOOKMARK_ID += 1
    return BOOKMARK_ID


def _apply_font(run, bold=False, size=Pt(14)):
    run.font.name = "Times New Roman"
    run.font.size = size
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold:
        run.font.bold = True


def _make_run_props(bold=False, size=Pt(14)):
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(size.pt * 2))
    rPr.append(sz)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '000000')
    rPr.append(color)
    if bold:
        b = OxmlElement('w:b')
        rPr.append(b)
    return rPr


def _add_p(doc, text, align=None, indent=Cm(1.25), hanging=None, bold=False, size=Pt(14)):
    p = doc.add_paragraph(text)
    p.paragraph_format.line_spacing = 1.5
    if align is not None:
        p.alignment = align
    if indent is not None:
        p.paragraph_format.first_line_indent = indent
    else:
        p.paragraph_format.first_line_indent = Cm(0)
    if hanging is not None:
        p.paragraph_format.left_indent = hanging
        p.paragraph_format.first_line_indent = -hanging
    for run in p.runs:
        _apply_font(run, bold=bold, size=size)
    return p


def _add_heading_p(doc, text, is_structural=False):
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    if is_structural:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Cm(1.25)
    for run in p.runs:
        _apply_font(run, bold=True)
    return p


def _add_subheading_p(doc, text):
    p = doc.add_paragraph(text, style="Heading 2")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    for run in p.runs:
        _apply_font(run, bold=True)
    return p


def _add_bookmark(doc, paragraph, name):
    bm_id = _next_bm_id()
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bm_id))
    start.set(qn('w:name'), name)
    paragraph._p.append(start)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bm_id))
    paragraph._p.append(end)


def _add_hyperlinks_to_paragraph(p, text: str):
    pattern = r'\[(\d+(?:[-,]\s*\d+)*)\]'
    parts = re.split(pattern, text)
    for i, part in enumerate(parts):
        if i % 2 == 0:
            if part:
                run = p.add_run(part)
                _apply_font(run)
        else:
            nums = re.findall(r'\d+', part)
            if nums:
                first = nums[0]
                hl = OxmlElement('w:hyperlink')
                hl.set(qn('w:anchor'), f'ref-{first}')
                hl.set(qn('w:history'), '1')
                r = OxmlElement('w:r')
                r.append(_make_run_props())
                t = OxmlElement('w:t')
                t.text = f'[{part}]'
                r.append(t)
                hl.append(r)
                p._p.append(hl)


def _add_text_paragraph(doc, text):
    pattern = r'\[(\d+(?:[-,]\s*\d+)*)\]'
    if re.search(pattern, text):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _add_hyperlinks_to_paragraph(p, text)
        return p
    return _add_p(doc, text)


def _is_separator_row(text):
    cleaned = text.replace("|", "").replace("-", "").replace(" ", "").replace(":", "")
    return len(cleaned) == 0


def _add_table(doc, rows: List[str]):
    data_rows = [r for r in rows if not _is_separator_row(r)]
    if not data_rows:
        return
    parsed = []
    for row in data_rows:
        cells = [c.strip() for c in row.strip("|").split("|")]
        parsed.append(cells)
    num_cols = max(len(r) for r in parsed)
    if num_cols == 0:
        return
    table = doc.add_table(rows=len(parsed), cols=num_cols)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row_data in enumerate(parsed):
        for j, cell_text in enumerate(row_data):
            if j >= num_cols:
                break
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(cell_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if i == 0:
                run.font.bold = True
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)


def _process_paragraphs(doc, paragraphs: List[str]):
    i = 0
    while i < len(paragraphs):
        text = paragraphs[i].strip()
        if not text:
            i += 1
            continue
        if text.startswith("|") and text.endswith("|"):
            rows = []
            while i < len(paragraphs):
                line = paragraphs[i].strip()
                if line.startswith("|") and line.endswith("|"):
                    rows.append(line)
                    i += 1
                else:
                    break
            _add_table(doc, rows)
        else:
            _add_text_paragraph(doc, text)
            i += 1


def _add_page_numbers(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    footer = section.footer
    footer.is_linked_to_previous = False

    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)

    r1 = p.add_run()
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(14)
    r1.font.color.rgb = RGBColor(0, 0, 0)
    fc1 = OxmlElement('w:fldChar')
    fc1.set(qn('w:fldCharType'), 'begin')
    r1._r.append(fc1)

    r2 = p.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = ' PAGE '
    r2._r.append(it)

    r3 = p.add_run()
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(14)
    r3.font.color.rgb = RGBColor(0, 0, 0)
    fc2 = OxmlElement('w:fldChar')
    fc2.set(qn('w:fldCharType'), 'end')
    r3._r.append(fc2)


def _add_toc(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)

    r1 = paragraph.add_run()
    fc1 = OxmlElement('w:fldChar')
    fc1.set(qn('w:fldCharType'), 'begin')
    r1._r.append(fc1)

    r2 = paragraph.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = ' TOC \\o "1-2" \\h '
    r2._r.append(it)

    r3 = paragraph.add_run()
    fc2 = OxmlElement('w:fldChar')
    fc2.set(qn('w:fldCharType'), 'separate')
    r3._r.append(fc2)

    r4 = paragraph.add_run()
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(10)
    r4.font.color.rgb = RGBColor(128, 128, 128)
    r4.text = "(обновите оглавление: правый клик -> Обновить поле)"

    r5 = paragraph.add_run()
    fc3 = OxmlElement('w:fldChar')
    fc3.set(qn('w:fldCharType'), 'end')
    r5._r.append(fc3)


def build_report(job_data: Dict[str, Any]) -> bytes:
    global BOOKMARK_ID
    BOOKMARK_ID = 0

    if os.path.exists(TEMPLATE_PATH):
        doc = Document(TEMPLATE_PATH)
    else:
        doc = Document()

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    try:
        hl = doc.styles["Hyperlink"]
    except KeyError:
        hl = doc.styles.add_style("Hyperlink", 2)
    hl.font.color.rgb = RGBColor(0, 0, 0)
    hl.font.underline = False

    _add_page_numbers(doc)
    sections_data = job_data.get("sections", {})

    _add_title_page(doc)
    doc.add_page_break()

    _add_abstract(doc, sections_data.get("abstract", {}))
    doc.add_page_break()

    _add_terms(doc, sections_data.get("terms", {}))
    doc.add_page_break()

    _add_abbreviations(doc, sections_data.get("abbreviations", {}))
    doc.add_page_break()

    _add_heading_p(doc, "СОДЕРЖАНИЕ", is_structural=True)
    _add_toc(doc)
    doc.add_page_break()

    _add_section(doc, "ВВЕДЕНИЕ", sections_data.get("introduction", {}))
    doc.add_page_break()

    for ch_key in ["chapter1", "chapter2", "chapter3"]:
        chapter = sections_data.get(ch_key, {})
        if chapter:
            _add_chapter(doc, chapter)
            doc.add_page_break()

    _add_section(doc, "ЗАКЛЮЧЕНИЕ", sections_data.get("conclusion", {}))
    doc.add_page_break()

    _add_references(doc, sections_data.get("references", {}))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_section(doc, title, data):
    if not data:
        return
    _add_heading_p(doc, title, is_structural=True)
    _process_paragraphs(doc, data.get("paragraphs", []))


def _add_chapter(doc, chapter):
    heading = chapter.get("heading", "ГЛАВА")
    _add_heading_p(doc, heading, is_structural=False)
    for sub in chapter.get("subsections", []):
        sub_heading = sub.get("heading", "")
        if sub_heading:
            _add_subheading_p(doc, sub_heading)
        _process_paragraphs(doc, sub.get("paragraphs", []))


def _add_abstract(doc, abstract):
    if not abstract:
        return
    _add_heading_p(doc, "РЕФЕРАТ", is_structural=True)
    volume = abstract.get("volume", "")
    if volume:
        _add_p(doc, volume, indent=Cm(0))
    keywords = abstract.get("keywords", [])
    if keywords:
        p = _add_p(doc, "", indent=Cm(0))
        run = p.add_run("Ключевые слова: ")
        _apply_font(run, bold=True)
        run = p.add_run(", ".join(k.upper() for k in keywords))
        _apply_font(run)
    text = abstract.get("text", "")
    if text:
        _add_p(doc, text)


def _add_references(doc, references):
    if not references:
        return
    _add_heading_p(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", is_structural=True)
    for idx, ref in enumerate(references.get("references", []), 1):
        p = _add_p(doc, ref, hanging=Cm(1.25))
        _add_bookmark(doc, p, f"ref-{idx}")


def _add_title_page(doc):
    p = doc.add_paragraph("Здесь будет титульник, листай ниже")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(200)
    for run in p.runs:
        _apply_font(run)


def _add_terms(doc, terms):
    if not terms:
        return
    _add_heading_p(doc, "ТЕРМИНЫ И ОПРЕДЕЛЕНИЯ", is_structural=True)
    opening = terms.get("opening", "В настоящем отчете о НИР применяют следующие термины с соответствующими определениями:")
    _add_p(doc, opening)
    for item in terms.get("items", []):
        term = item.get("term", "")
        definition = item.get("definition", "")
        text = f"{term} — {definition}" if definition else term
        _add_p(doc, text)


def _add_abbreviations(doc, abbreviations):
    if not abbreviations:
        return
    _add_heading_p(doc, "ПЕРЕЧЕНЬ СОКРАЩЕНИЙ И ОБОЗНАЧЕНИЙ", is_structural=True)
    opening = abbreviations.get("opening", "В настоящем отчете о НИР применяют следующие сокращения и обозначения:")
    _add_p(doc, opening)
    for item in abbreviations.get("items", []):
        abbr = item.get("abbr", "")
        full = item.get("full", "")
        text = f"{abbr} — {full}" if full else abbr
        _add_p(doc, text)
