from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml

_FIG_COUNTER = 0
_TBL_COUNTER = 0
_LST_COUNTER = 0


def _next_fig_num():
    global _FIG_COUNTER
    _FIG_COUNTER += 1
    return _FIG_COUNTER


def _next_tbl_num():
    global _TBL_COUNTER
    _TBL_COUNTER += 1
    return _TBL_COUNTER


def _next_lst_num():
    global _LST_COUNTER
    _LST_COUNTER += 1
    return _LST_COUNTER


def reset_counters():
    global _FIG_COUNTER, _TBL_COUNTER, _LST_COUNTER
    _FIG_COUNTER = 0
    _TBL_COUNTER = 0
    _LST_COUNTER = 0


def make_paragraph(doc, text, font_name="Times New Roman", font_size=14, bold=False,
                   italic=False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                   space_before=0, space_after=0, line_spacing=1.5,
                   first_line_indent=None, keep_with_next=False,
                   page_break_before=False, all_caps=False, style=None):
    if isinstance(text, (list, tuple)):
        text = text[0] if text else ""
    text = str(text)
    if all_caps:
        text = text.upper()
    if style:
        p = doc.add_paragraph(style=style)
        if first_line_indent is not None:
            p.paragraph_format.first_line_indent = Cm(first_line_indent)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        if line_spacing:
            p.paragraph_format.line_spacing = line_spacing
        if alignment is not None:
            p.alignment = alignment
        if page_break_before:
            p.paragraph_format.page_break_before = True
        if keep_with_next:
            p.paragraph_format.keep_with_next = True
    else:
        p = doc.add_paragraph()
        if page_break_before:
            p.paragraph_format.page_break_before = True
        if first_line_indent is not None:
            p.paragraph_format.first_line_indent = Cm(first_line_indent)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        if line_spacing:
            p.paragraph_format.line_spacing = line_spacing
        p.alignment = alignment
        if keep_with_next:
            p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    run.italic = italic
    rpr = run._r.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="{font_name}"/>')
        rpr.append(rFonts)
    else:
        rFonts.set(qn('w:eastAsia'), font_name)
    return p


def make_table(doc, caption_text, headers, rows, font_size=12):
    tbl_num = _next_tbl_num()
    make_paragraph(doc, f"Таблица {tbl_num} — {caption_text}", font_size=14,
                   alignment=WD_ALIGN_PARAGRAPH.LEFT,
                   first_line_indent=0, space_before=3, space_after=0,
                   line_spacing=1.0)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '  <w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '  <w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '  <w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '</w:tblBorders>'
    ))
    t._tbl.tblPr.append(parse_xml(f'<w:tblLook {nsdecls("w")} w:firstRow="1"/>'))
    hdr = t.rows[0]
    hdr._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")} w:val="true"/>'))
    hdr._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")} w:val="true"/>'))
    for ci, h in enumerate(headers):
        cell = hdr.cells[ci]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.bold = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
    for ri, row in enumerate(rows):
        t.rows[ri + 1]._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")} w:val="true"/>'))
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            run.font.name = "Times New Roman"
            run.font.size = Pt(font_size)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
    make_paragraph(doc, "", font_size=14, space_before=0, space_after=8,
                   first_line_indent=0, line_spacing=1.0)


def make_figure(doc, caption_text, comment=None):
    fig_num = _next_fig_num()
    if comment:
        make_paragraph(doc, comment, font_size=12, italic=True,
                       alignment=WD_ALIGN_PARAGRAPH.CENTER,
                       first_line_indent=0, space_before=6, space_after=0,
                       line_spacing=1.5)
    make_paragraph(doc, f"Рисунок {fig_num} — {caption_text}", font_size=12,
                   alignment=WD_ALIGN_PARAGRAPH.CENTER,
                   first_line_indent=1.25, space_before=3, space_after=6,
                   line_spacing=1.5)


def make_code_block(doc, code_text, caption_text=None):
    if caption_text:
        # Auto-number if caption doesn't already contain "Листинг N"
        if "Листинг" not in caption_text:
            lst_num = _next_lst_num()
            caption_text = f"Листинг {lst_num} — {caption_text}"
        else:
            _next_lst_num()  # advance counter even if caption pre-formatted
        make_paragraph(doc, caption_text, font_size=12, bold=True,
                       alignment=WD_ALIGN_PARAGRAPH.LEFT,
                       first_line_indent=0, space_before=6, space_after=0,
                       line_spacing=1.0)
    lines = code_text.strip().split("\n")[:25]
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '  <w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '  <w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '  <w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '</w:tblBorders>'
    )
    tbl._tbl.tblPr.append(borders)
    tbl._tbl.tblPr.append(parse_xml(f'<w:tblLook {nsdecls("w")} w:firstRow="1"/>'))
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")} w:val="true"/>'))
    cell = tbl.rows[0].cells[0]
    for i, line in enumerate(lines):
        if i == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(10)
    make_paragraph(doc, "", font_size=10, space_before=0, space_after=8,
                   first_line_indent=0, line_spacing=1.0)


def add_reference(doc, text):
    return make_paragraph(doc, text, first_line_indent=0,
                          space_before=0, space_after=3, line_spacing=1.5)


def add_page_numbers(doc):
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fld = OxmlElement('w:fldChar')
        fld.set(qn('w:fldCharType'), 'begin')
        run._r.append(fld)
        run2 = p.add_run()
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = 'PAGE'
        run2._r.append(instr)
        run3 = p.add_run()
        fld2 = OxmlElement('w:fldChar')
        fld2.set(qn('w:fldCharType'), 'end')
        run3._r.append(fld2)


def add_definition_table(doc, items):
    if not items:
        return
    table = doc.add_table(rows=len(items), cols=2)
    for i, (left, right) in enumerate(items):
        trPr = table.rows[i]._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")} w:val="true"/>'))
        for ci, val in enumerate([left, f"— {right}"]):
            cell = table.rows[i].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            if ci == 0:
                run.bold = True
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.5
        table.rows[i].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
